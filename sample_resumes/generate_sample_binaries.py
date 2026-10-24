import os
import fitz  # PyMuPDF
import docx  # python-docx

output_dir = os.path.dirname(os.path.abspath(__file__))

# 1. Generate Alex Mercer (PDF)
alex_text = """Alex Mercer
Email: alex.mercer.dev@example.com | Phone: +1 (555) 382-9102
Location: San Francisco, CA | LinkedIn: linkedin.com/in/alex-mercer | GitHub: github.com/alexmercer

PROFESSIONAL SUMMARY
Senior Full Stack & AI Systems Engineer with 5+ years of experience designing and deploying scalable web architectures, microservices, and automated NLP systems. Proven track record leading distributed systems with Python, FastAPI, React, Next.js, and PostgreSQL in containerized Kubernetes and Docker environments.

CORE TECHNICAL SKILLS
- Programming: Python, TypeScript, JavaScript, SQL, HTML5, CSS3
- Frontend: React, Next.js, Tailwind CSS, Redux, Webpack
- Backend: FastAPI, Node.js, Express, REST APIs, GraphQL
- Databases: PostgreSQL, Redis, Supabase, MongoDB
- Cloud & DevOps: Docker, Kubernetes, AWS (S3, EC2, ECS), CI/CD (GitHub Actions), Linux

PROFESSIONAL EXPERIENCE
Senior Full Stack AI Engineer | CloudSphere Solutions | 2022 - Present
- Architected high-throughput microservices using Python and FastAPI serving 1M+ daily active requests with sub-100ms p95 latency.
- Designed reactive client dashboard in React & Next.js connected to PostgreSQL RDS cluster.
- Implemented real-time resume parsing and embedding generation using PyMuPDF and Scikit-learn cosine similarity.
- Containerized application fleet with Docker and integrated automated CI/CD deployment pipelines on AWS ECS.

Full Stack Software Developer | InnovateTech Labs | 2019 - 2022
- Developed responsive web applications using React, TypeScript, and Tailwind CSS.
- Built performant RESTful APIs using Python and PostgreSQL for customer analytics.
- Reduced database query times by 40% through indexing and query optimization on PostgreSQL.

EDUCATION
Bachelor of Science in Computer Science
University of Washington, Seattle, WA (2015 - 2019)

PROJECTS
- AI Candidate Matcher: Full stack application built with FastAPI, Next.js, and scikit-learn for intelligent document scoring.
- CloudMetrics: Real-time telemetry dashboard using React, WebSockets, and Redis.
"""

doc_alex = fitz.open()
page_alex = doc_alex.new_page()
page_alex.insert_text(fitz.Point(40, 50), alex_text, fontsize=10)
pdf_alex_path = os.path.join(output_dir, "alex_mercer_lead_ai.pdf")
doc_alex.save(pdf_alex_path)
doc_alex.close()
print(f"Generated PDF: {pdf_alex_path}")

# 2. Generate Sarah Chen (DOCX)
doc_sarah = docx.Document()
doc_sarah.add_heading("Sarah Chen", level=0)
doc_sarah.add_paragraph("Email: sarah.chen.tech@example.com | Phone: +1 (555) 749-2184")
doc_sarah.add_paragraph("Location: Austin, TX | LinkedIn: linkedin.com/in/sarahchen-swe")

doc_sarah.add_heading("Professional Summary", level=1)
doc_sarah.add_paragraph(
    "Accomplished Full Stack Web Developer with 4 years of experience building modern user-facing applications. "
    "Strong focus on clean code, TypeScript, React ecosystem, and Python backend services."
)

doc_sarah.add_heading("Technical Skills", level=1)
doc_sarah.add_paragraph("Languages: Python, JavaScript, TypeScript, SQL")
doc_sarah.add_paragraph("Frontend: React, Next.js, Tailwind CSS, HTML5, CSS3, Vite")
doc_sarah.add_paragraph("Backend: FastAPI, Node.js, REST APIs")
doc_sarah.add_paragraph("Databases: PostgreSQL, SQLite")
doc_sarah.add_paragraph("Tooling: Git, GitHub, Postman, Jest")

doc_sarah.add_heading("Work Experience", level=1)
doc_sarah.add_paragraph("Software Engineer | NexaDigital (2021 - Present)")
doc_sarah.add_paragraph("- Implemented REST APIs in Python using FastAPI with PostgreSQL persistence layer.")
doc_sarah.add_paragraph("- Built modular UI component library in React with responsive Tailwind CSS styling.")
doc_sarah.add_paragraph("- Collaborated with product designers to implement pixel-perfect user journeys in Next.js.")

doc_sarah.add_heading("Education", level=1)
doc_sarah.add_paragraph("Master of Science in Software Systems - Georgia Institute of Technology (2018 - 2020)")
doc_sarah.add_paragraph("Bachelor of Science in Information Systems - University of Texas at Austin (2014 - 2018)")

docx_sarah_path = os.path.join(output_dir, "sarah_chen_fullstack.docx")
doc_sarah.save(docx_sarah_path)
print(f"Generated DOCX: {docx_sarah_path}")

# 3. Generate David Patel (PDF)
david_text = """David Patel
Email: david.patel.cloud@example.com | Phone: +1 (555) 492-8172
Location: San Jose, CA | LinkedIn: linkedin.com/in/davidpatel | GitHub: github.com/davidpatel

OBJECTIVE
Motivated software engineer with 2 years of experience looking to contribute to innovative web development teams.

SKILLS
- Programming: Python, JavaScript, HTML, CSS
- Frameworks: React, Express, Node.js
- Databases: MongoDB, SQLite
- Tools: Docker, Git, Linux, Postman

EXPERIENCE
Associate Software Developer | DataWave Tech | 2022 - Present
- Built client-side interfaces using React and CSS.
- Developed backend microservices using Python and Flask for internal data transformations.
- Managed containerized environments with basic Docker commands.

EDUCATION
Bachelor of Engineering in Information Technology
Pune University (2018 - 2022)

PROJECTS
- TaskMaster: Simple task board application using React and MongoDB.
- WeatherCast: CLI tool written in Python to fetch forecast data from open APIs.
"""

doc_david = fitz.open()
page_david = doc_david.new_page()
page_david.insert_text(fitz.Point(40, 50), david_text, fontsize=10)
pdf_david_path = os.path.join(output_dir, "david_patel_junior_dev.pdf")
doc_david.save(pdf_david_path)
doc_david.close()
print(f"Generated PDF: {pdf_david_path}")
