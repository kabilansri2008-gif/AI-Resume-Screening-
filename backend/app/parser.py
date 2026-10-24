import io
import re
from typing import Dict, List, Optional, Tuple, Any

try:
    import fitz  # PyMuPDF
except ImportError:
    fitz = None

try:
    import docx
except ImportError:
    docx = None


COMMON_SKILLS_DB = {
    # Programming Languages
    "python", "javascript", "typescript", "java", "c++", "c#", "c", "ruby", "php", "swift", "kotlin", "go", "golang", "rust", "scala", "r", "dart",
    # Frontend
    "react", "react.js", "next.js", "vue", "vue.js", "angular", "svelte", "html", "html5", "css", "css3", "tailwind", "tailwind css", "sass", "bootstrap", "redux", "zustand", "graphql", "rest api", "webpack", "vite",
    # Backend
    "node.js", "nodejs", "express", "express.js", "fastapi", "django", "flask", "spring boot", "asp.net", "rails", "nestjs", "gin", "grpc",
    # Databases & Storage
    "postgresql", "postgres", "mysql", "mongodb", "sqlite", "redis", "elasticsearch", "dynamodb", "supabase", "firebase", "cassandra", "oracle", "mariadb",
    # Cloud & DevOps
    "aws", "amazon web services", "azure", "gcp", "google cloud", "docker", "kubernetes", "terraform", "ansible", "ci/cd", "jenkins", "github actions", "gitlab ci", "linux", "nginx", "helm",
    # AI / ML / Data Science
    "machine learning", "deep learning", "nlp", "natural language processing", "computer vision", "llm", "pytorch", "tensorflow", "keras", "scikit-learn", "pandas", "numpy", "matplotlib", "seaborn", "huggingface", "langchain", "llamaindex", "openai", "transformers", "xgboost",
    # Testing & Tooling
    "jest", "cypress", "playwright", "pytest", "unit testing", "git", "github", "gitlab", "jira", "postman", "agile", "scrum", "microservices"
}

SYNONYM_MAP = {
    "reactjs": "react",
    "react.js": "react",
    "nodejs": "node.js",
    "node": "node.js",
    "ts": "typescript",
    "js": "javascript",
    "golang": "go",
    "py": "python",
    "k8s": "kubernetes",
    "postgres": "postgresql",
    "tf": "tensorflow",
    "sk-learn": "scikit-learn",
    "sklearn": "scikit-learn",
    "tailwind": "tailwind css",
    "next": "next.js",
    "nextjs": "next.js",
    "vuejs": "vue",
    "vue.js": "vue",
    "aws": "amazon web services",
    "gcp": "google cloud",
    "rest": "rest api",
    "restful": "rest api",
    "ml": "machine learning",
    "ai": "artificial intelligence"
}


def normalize_skill(skill: str) -> str:
    cleaned = skill.strip().lower()
    cleaned = re.sub(r'[\s\-_]+', ' ', cleaned)
    return SYNONYM_MAP.get(cleaned, cleaned)


def extract_text_from_pdf(file_bytes: bytes) -> Tuple[str, List[str]]:
    """Extracts text from PDF using PyMuPDF (fitz) and gathers parsing warnings."""
    warnings: List[str] = []
    if not fitz:
        raise RuntimeError("PyMuPDF (fitz) is not installed in the Python environment.")

    try:
        doc = fitz.open(stream=file_bytes, filetype="pdf")
    except Exception as e:
        raise ValueError(f"Corrupted or invalid PDF format: {str(e)}")

    if doc.is_encrypted:
        raise ValueError("Resume file is password-protected. Please upload an unprotected PDF.")

    if doc.page_count == 0:
        raise ValueError("The uploaded PDF has 0 pages (empty document).")

    full_text = []
    has_text = False

    for page_num in range(doc.page_count):
        page = doc.load_page(page_num)
        page_text = page.get_text() or ""
        if page_text.strip():
            has_text = True
            full_text.append(page_text)

    combined_text = "\n".join(full_text).strip()

    if not has_text or len(combined_text) < 40:
        warnings.append(
            "Very little text was extracted. The PDF might be a scanned image or photo. OCR fallback recommendation applied."
        )

    return combined_text, warnings


def extract_text_from_docx(file_bytes: bytes) -> Tuple[str, List[str]]:
    """Extracts text from DOCX using python-docx and gathers parsing warnings."""
    warnings: List[str] = []
    if not docx:
        raise RuntimeError("python-docx is not installed in the Python environment.")

    try:
        doc = docx.Document(io.BytesIO(file_bytes))
    except Exception as e:
        raise ValueError(f"Corrupted or invalid DOCX format: {str(e)}")

    text_parts = []
    for p in doc.paragraphs:
        if p.text.strip():
            text_parts.append(p.text.strip())

    for table in doc.tables:
        for row in table.rows:
            row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_texts:
                text_parts.append(" | ".join(row_texts))

    combined = "\n".join(text_parts).strip()
    if not combined or len(combined) < 40:
        warnings.append("Document appears empty or contains non-extractable elements.")

    return combined, warnings


def extract_contact_info(text: str) -> Dict[str, Optional[str]]:
    email_pattern = r'[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+'
    phone_pattern = r'(\+?\d{1,3}[-.\s]?)?\(?\d{3}\)?[-.\s]?\d{3}[-.\s]?\d{4}'
    linkedin_pattern = r'linkedin\.com/in/[a-zA-Z0-9_-]+'
    github_pattern = r'github\.com/[a-zA-Z0-9_-]+'

    emails = re.findall(email_pattern, text)
    phones = re.findall(phone_pattern, text)
    linkedins = re.findall(linkedin_pattern, text, re.IGNORECASE)
    githubs = re.findall(github_pattern, text, re.IGNORECASE)

    # Name extraction heuristic: check the top 4 non-empty lines
    lines = [line.strip() for line in text.split('\n') if line.strip()]
    candidate_name = None
    for line in lines[:4]:
        # Filter out common headings or contact lines
        if any(h in line.lower() for h in ['resume', 'curriculum', 'cv', '@', 'http', 'phone', 'page', 'email']):
            continue
        words = line.split()
        if 1 <= len(words) <= 4 and all(re.match(r'^[A-Za-z.\'-]+$', w) for w in words):
            candidate_name = line
            break

    return {
        "name": candidate_name or "Candidate",
        "email": emails[0] if emails else None,
        "phone": phones[0] if phones else None,
        "linkedin": f"https://{linkedins[0]}" if linkedins else None,
        "github": f"https://{githubs[0]}" if githubs else None,
    }


def extract_skills_from_text(text: str) -> List[str]:
    lowered = text.lower()
    matched_skills = set()

    for skill in COMMON_SKILLS_DB:
        # Pattern ensures word boundaries or punctuation boundaries (e.g. c++, c#, node.js)
        escaped = re.escape(skill)
        pattern = rf'(?<![a-zA-Z0-9]){escaped}(?![a-zA-Z0-9])'
        if re.search(pattern, lowered):
            matched_skills.add(normalize_skill(skill))

    return sorted(list(matched_skills))


def extract_experience_years(text: str) -> float:
    """Estimates years of experience from patterns like '5+ years', '3 years of experience'."""
    pattern = r'(\d{1,2}(?:\.\d+)?)\+?\s*(?:years?|yrs?)\b(?:\s+of\s+experience)?'
    matches = re.findall(pattern, text, re.IGNORECASE)
    if matches:
        try:
            years = [float(m) for m in matches if float(m) <= 45]
            if years:
                return max(years)
        except ValueError:
            pass
    return 0.0


def extract_education(text: str) -> List[str]:
    edu_keywords = [
        "bachelor", "master", "ph.d", "phd", "b.s.", "m.s.", "b.tech", "m.tech",
        "b.e.", "m.e.", "bca", "mca", "b.sc", "m.sc", "diploma", "degree"
    ]
    results = []
    lines = text.split('\n')
    for line in lines:
        l_lower = line.lower()
        if any(k in l_lower for k in edu_keywords):
            cleaned = line.strip()
            if len(cleaned) < 120 and cleaned not in results:
                results.append(cleaned)
    return results[:4]


def extract_projects(text: str) -> List[str]:
    lines = text.split('\n')
    in_project_section = False
    project_snippets = []

    for line in lines:
        cleaned = line.strip()
        lower = cleaned.lower()
        if any(header in lower for header in ['projects', 'key projects', 'academic projects', 'featured work']):
            in_project_section = True
            continue
        if in_project_section:
            if any(header in lower for header in ['education', 'skills', 'experience', 'certifications', 'hobbies']):
                in_project_section = False
                continue
            if cleaned.startswith(('•', '-', '*', '1.', '2.', '3.')) or (len(cleaned) > 20 and len(cleaned) < 160):
                project_snippets.append(cleaned.lstrip('•-* 0123456789.'))

    return project_snippets[:6]


def parse_resume(file_bytes: bytes, filename: str) -> Dict[str, Any]:
    ext = filename.lower().split('.')[-1]
    warnings: List[str] = []

    if ext == 'pdf':
        text, parse_warnings = extract_text_from_pdf(file_bytes)
        warnings.extend(parse_warnings)
    elif ext in ['docx', 'doc']:
        text, parse_warnings = extract_text_from_docx(file_bytes)
        warnings.extend(parse_warnings)
    elif ext in ['txt', 'md']:
        text = file_bytes.decode('utf-8', errors='ignore')
    else:
        raise ValueError(f"Unsupported file format '.{ext}'. Supported formats are PDF, DOCX, and TXT.")

    contact = extract_contact_info(text)
    if not contact["email"]:
        warnings.append("No valid email address found in the resume header.")
    if not contact["phone"]:
        warnings.append("No phone number detected in the contact details.")

    skills = extract_skills_from_text(text)
    exp_years = extract_experience_years(text)
    education = extract_education(text)
    projects = extract_projects(text)

    return {
        "filename": filename,
        "name": contact["name"],
        "email": contact["email"] or "Not specified",
        "phone": contact["phone"] or "Not specified",
        "linkedin": contact["linkedin"],
        "github": contact["github"],
        "skills": skills,
        "experience_years": exp_years,
        "education": education,
        "projects": projects,
        "raw_text": text,
        "word_count": len(text.split()),
        "warnings": warnings,
        "confidence_score": 95 if len(warnings) == 0 else (75 if len(warnings) == 1 else 55)
    }
